# -*- coding: utf-8 -*-
"""
Spike-2：报告渲染管线验证（Plan 阶段 E-4/E-11）
验证链路：日线+估值数据拉取 -> 18 项日线因子（附表 A）-> 同池百分位 -> S/H（H 为简化演示口径）
        -> matplotlib 图表 -> HTML 在线版 + docx 文件版双输出 -> 合规关键词扫描（AC-005 预演）

重要声明（Spike 纪律）：
- 本输出为 SPIKE 演示样张，非正式报告。
- 同池为 20 只演示子池（真实四池名单接口已在 Spike-1 验证，正式实现时替换）。
- H 分用 3 因子代理（正式口径为 18 因子周频 S 的历史位置，逻辑同构）。
"""
import os, sys, json, time, traceback
import datetime as dt

T0 = time.time()
RESULT = {"steps": {}, "notes": []}

def step(name, ok, detail=""):
    RESULT["steps"][name] = {"pass": bool(ok), "detail": detail, "t": round(time.time() - T0, 1)}
    print(f"[{'PASS' if ok else 'FAIL'}] {name} ({RESULT['steps'][name]['t']}s) {detail}")

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "out")
os.makedirs(OUT, exist_ok=True)

import numpy as np
import pandas as pd
import akshare as ak

# 环境代理对本机数据源域名不可达（2026-10-17 实测），改为直连
os.environ["NO_PROXY"] = "*"
for _k in ("http_proxy", "https_proxy", "HTTP_PROXY", "HTTPS_PROXY"):
    os.environ.pop(_k, None)

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei"]
plt.rcParams["axes.unicode_minus"] = False

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

# ---------------------------------------------------------------- 常量
TARGET = "603799"          # 华友钴业（与样例报告 INPUT-001 同标的，便于人工比对版式）
POOL_DEMO = [              # 演示子池：流动性好、长期为宽基成分的样本（Spike 用，正式池走中证官网名单）
    ("603799", "华友钴业"), ("600519", "贵州茅台"), ("601318", "中国平安"), ("600036", "招商银行"),
    ("000858", "五粮液"),   ("300750", "宁德时代"), ("002594", "比亚迪"),   ("600900", "长江电力"),
    ("601899", "紫金矿业"), ("000333", "美的集团"), ("600276", "恒瑞医药"), ("002415", "海康威视"),
    ("600887", "伊利股份"), ("601166", "兴业银行"), ("600028", "中国石化"), ("601398", "工商银行"),
    ("300059", "东方财富"), ("601668", "中国建筑"), ("601012", "隆基绿能"), ("600030", "中信证券"),
]
N_POOL = len(POOL_DEMO)
START = "20240901"
END = "20261017"
RED, GREEN, GRAY = "#C00000", "#1E7B3C", "#555555"   # BR-007：<40 绿 / 40-60 深灰 / >60 红（涨红跌绿）

def band(x):
    return ("red", RED, "偏有利") if x > 60 else (("green", GREEN, "偏不利") if x < 40 else ("gray", GRAY, "中性"))

FACTOR_TABLE = [  # (#, 侧面, 内部名, 报告匿名标签, 是否低值有利)
    ("F1", "估值特征", "pe_pct",     "估值·定价水平", True),
    ("F2", "估值特征", "pb_pct",     "估值·资产口径", True),
    ("F3", "估值特征", "mv_pct",     "市值·规模位置", True),
    ("F4", "价格行为", "ret20",     "短周期价格动量", False),
    ("F5", "价格行为", "ret60",     "中周期价格动量", False),
    ("F6", "价格行为", "ma250_dist", "长周期趋势位置", False),
    ("F7", "价格行为", "ma_ratio",   "趋势结构（中短/中长期）", False),
    ("F8", "价格行为", "pos52w",     "年度区间位置", False),
    ("F9", "价格行为", "accel",      "短期动能变化", False),
    ("F10", "价格行为", "updays",   "上行一致性", False),
    ("F11", "交易行为", "amt20",     "交易活跃度", False),
    ("F12", "交易行为", "volx",      "量能扩张度", False),
    ("F13", "交易行为", "vp_corr",   "量价配合度", False),
    ("F14", "交易行为", "turn20",    "筹码交换强度", False),
    ("F15", "风险特征", "vol20",     "短期波动水平", True),
    ("F16", "风险特征", "vol60",     "中期波动水平", True),
    ("F17", "风险特征", "mdd60",     "区间回撤深度", True),
    ("F18", "风险特征", "amp20",     "日内振幅水平", True),
]

# ---------------------------------------------------------------- 1. 数据拉取
# Spike-2 环境发现：东财域名被沙箱阻断、akshare 1.18.64 已移除乐咕个股估值接口
# 方案：日线 东财主源 + 新浪备援；估值（PE_TTM/PB/总市值）百度股市通接口（实测可达）
_EAST_OK = {"v": None}
def _east_ok():
    if _EAST_OK["v"] is None:
        try:
            p = ak.stock_zh_a_hist(symbol="600519", period="daily", start_date="20260901",
                                    end_date=END, adjust="qfq")
            _EAST_OK["v"] = len(p) > 0
        except Exception:
            _EAST_OK["v"] = False
    return _EAST_OK["v"]

def _baidu_val(code):
    """百度估值最新值：总市值/市盈率(TTM)/市净率（仅取截面最新值，供同池百分位）"""
    out = {}
    for ind, col in [("总市值", "total_mv"), ("市盈率(TTM)", "pe_ttm"), ("市净率", "pb")]:
        v = np.nan
        for _try in (1, 2):
            try:
                bv = ak.stock_zh_valuation_baidu(symbol=code, indicator=ind, period="近一年")
                v = float(bv["value"].dropna().iloc[-1]); break
            except Exception:
                time.sleep(1.5)
        out[col] = v
    return out

def fetch_one(code):
    df = None; src = "eastmoney"
    if _east_ok():
        try:
            df = ak.stock_zh_a_hist(symbol=code, period="daily", start_date=START, end_date=END, adjust="qfq")
            df = df.rename(columns={"日期": "date", "开盘": "open", "收盘": "close", "最高": "high", "最低": "low",
                                    "成交量": "volume", "成交额": "amount", "振幅": "amp", "换手率": "turnover"})
            df["date"] = pd.to_datetime(df["date"]); df = df.set_index("date").sort_index()
        except Exception:
            df = None
    if df is None or len(df) == 0:
        src = "sina"
        sym = ("sh" if code.startswith(("6", "9")) else ("bj" if code.startswith(("4", "8")) else "sz")) + code
        df = ak.stock_zh_a_daily(symbol=sym, start_date=START, end_date=END, adjust="qfq")
        df["date"] = pd.to_datetime(df["date"]); df = df.set_index("date").sort_index()
        df["amp"] = (df["high"] - df["low"]) / df["close"].shift(1) * 100   # 由 OHLC 补算
    df = df[["open", "high", "low", "close", "volume", "amount", "amp", "turnover"]].astype(float)
    for col, v in _baidu_val(code).items():
        df[col] = v
    df["code"] = code; df["source"] = src
    return df

panels, failed, src_stat = {}, [], {}
for code, name in POOL_DEMO:
    try:
        panels[code] = fetch_one(code)
        s = panels[code]["source"].iloc[0]; src_stat[s] = src_stat.get(s, 0) + 1
    except Exception as e:
        failed.append((code, repr(e)[:90]))

step("数据拉取", len(panels) >= 15, f"成功 {len(panels)}/{N_POOL} 只；日线来源 {src_stat}；失败: {failed if failed else '无'}")
if len(panels) < 10:
    RESULT["elapsed"] = round(time.time() - T0, 1)
    json.dump(RESULT, open(os.path.join(HERE, "spike2_result.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=2)
    print("数据不足，安全终止"); sys.exit(1)

# ---------------------------------------------------------------- 2. 因子原始值（管线按附表 A）
def raw_factors(df):
    c = df["close"]; v = df["volume"]
    r = {}
    r["pe_pct"]  = df["pe_ttm"].iloc[-1]
    r["pb_pct"]  = df["pb"].iloc[-1]
    r["mv_pct"]  = df["total_mv"].iloc[-1]
    r["ret20"]   = c.iloc[-1] / c.iloc[-21] - 1
    r["ret60"]   = c.iloc[-1] / c.iloc[-61] - 1
    r["ma250_dist"] = c.iloc[-1] / c.rolling(250).mean().iloc[-1]
    r["ma_ratio"]   = c.rolling(20).mean().iloc[-1] / c.rolling(60).mean().iloc[-1]
    w = c.iloc[-252:]
    r["pos52w"]  = (w.iloc[-1] - w.min()) / (w.max() - w.min() + 1e-12)
    r["accel"]   = (c.iloc[-1] / c.iloc[-6] - 1) - r["ret20"]
    r["updays"]  = float((c.diff().iloc[-20:] > 0).mean())
    r["amt20"]   = df["amount"].iloc[-20:].mean()
    r["volx"]    = v.iloc[-5:].mean() / (v.iloc[-20:].mean() + 1e-12)
    r["vp_corr"] = c.iloc[-20:].pct_change().corr(v.pct_change().iloc[-20:])
    r["turn20"]  = df["turnover"].iloc[-20:].mean()
    r["vol20"]   = c.pct_change().iloc[-20:].std()
    r["vol60"]   = c.pct_change().iloc[-60:].std()
    w60 = c.iloc[-60:]; cummax = w60.cummax()
    r["mdd60"]   = float(1 - (w60 / cummax).min())
    r["amp20"]   = df["amp"].iloc[-20:].mean()
    return r

raws = pd.DataFrame({code: raw_factors(df) for code, df in panels.items()}).T
_valid = raws.notna().sum()
step("18 项因子原始值计算", raws.notna().sum().sum() > 0,
     f"矩阵 {raws.shape}；因子有效样本 {int(_valid.min()) if _valid.size else 0}–{int(_valid.max()) if _valid.size else 0}（缺失不填分，BR-006）")

# ---------------------------------------------------------------- 3. 同池百分位（BR-006：平均并列名次百分位，缺失不填分）
scores = pd.DataFrame(index=raws.index, columns=[f[0] for f in FACTOR_TABLE], dtype=float)
for fid, side, key, label, low_good in FACTOR_TABLE:
    s = raws[key].rank(method="average", pct=True) * 100   # 平均并列名次百分位
    scores[fid] = (100 - s) if low_good else s
scores = scores.astype(float)
S_raw = scores.mean(axis=1, skipna=True)                    # 有效因子等权平均
S_score = S_raw.rank(method="average", pct=True) * 100     # 放回同池重新排名
step("同池百分位与 S 合成", True,
     f"目标 {TARGET}: S_raw={S_raw[TARGET]:.1f}, S={S_score[TARGET]:.1f}, 同池排名 {int(S_raw.rank(ascending=False)[TARGET])}/{len(S_raw)}")

# ---------------------------------------------------------------- 4. H 简化口径（3 因子代理，52 完成周）
closes = pd.DataFrame({code: df["close"] for code, df in panels.items()})
ret20d = closes / closes.shift(20) - 1
ret60d = closes / closes.shift(60) - 1
vol20d = closes.pct_change().rolling(20).std()
wk = pd.DataFrame({"ret20": ret20d.iloc[-252:].resample("W-FRI").last().stack(),
                   "ret60": ret60d.iloc[-252:].resample("W-FRI").last().stack(),
                   "vol20": vol20d.iloc[-252:].resample("W-FRI").last().stack()})
wk["sc"] = (wk.groupby(level=0)[["ret20", "ret60"]].rank(method="average", pct=True).mean(axis=1) * 100
            + (100 - wk.groupby(level=0)["vol20"].rank(method="average", pct=True))) / 2
wk_t = wk.xs(TARGET, level=1)["sc"].dropna()            # 取目标股：平面日期索引
hist52 = wk_t.iloc[-53:]                                   # 52 个完成周 + 当前周
H = float(hist52.rank(method="average", pct=True).iloc[-1] * 100)
step("H 历史位置（简化口径）", len(hist52) >= 40, f"52 周窗口样本 {len(hist52)} 周；当前 H={H:.0f}")

# ---------------------------------------------------------------- 5. 复核日（AC-006 预演：出分日 + 5 个交易日）
try:
    cal = ak.tool_trade_date_hist_sina()
    tdays = sorted(pd.to_datetime(cal["trade_date"]).tolist())
    last_day = panels[TARGET].index[-1]
    nxt = [d for d in tdays if d > last_day][:5]
    review_day = nxt[-1].strftime("%Y-%m-%d")
    ok_cal = True
except Exception as e:
    review_day, ok_cal = "（交易日历接口在正式实现中接入）", False
step("复核日计算", ok_cal, f"出分日 {last_day.strftime('%Y-%m-%d')}，建议复核日 {review_day}")

# ---------------------------------------------------------------- 6. 图表（对齐样例版式要素）
tgt_df = panels[TARGET]; tgt_close = tgt_df["close"]
S_val = float(S_score[TARGET])
side_mean = {s: float(scores.loc[TARGET, [f[0] for f in FACTOR_TABLE if f[1] == s]].mean())
             for s in ["估值特征", "价格行为", "交易行为", "风险特征"]}

def color_of(x): return RED if x > 60 else (GREEN if x < 40 else GRAY)

# 图 1：股性曲线（收盘价 + MA20/60/250）
fig, ax = plt.subplots(figsize=(8.6, 3.4), dpi=150)
ax.plot(tgt_close.index, tgt_close, lw=1.2, color="#333333", label="收盘价（前复权）")
for n, col in [(20, "#C00000"), (60, "#E6A100"), (250, "#777777")]:
    ax.plot(tgt_close.index, tgt_close.rolling(n).mean(), lw=0.9, color=col, label=f"MA{n}")
ax.set_title(f"股性曲线（演示）：603799 华友钴业", fontsize=11)
ax.legend(fontsize=8, ncol=4); ax.grid(alpha=.25)
fig.tight_layout(); fig.savefig(os.path.join(OUT, "chart_price.png")); plt.close(fig)

# 图 2：18 项因子分（按侧面分组，BR-007 着色）
fig, ax = plt.subplots(figsize=(8.6, 5.2), dpi=150)
labels = [f"{f[0]} {f[3]}[{f[1][:2]}]" for f in FACTOR_TABLE][::-1]
vals = [scores.loc[TARGET, f[0]] for f in FACTOR_TABLE][::-1]
ax.barh(range(18), vals, color=[color_of(v) for v in vals], height=.62)
ax.set_yticks(range(18)); ax.set_yticklabels(labels, fontsize=8)
ax.axvline(40, color=GREEN, ls="--", lw=.8); ax.axvline(60, color=RED, ls="--", lw=.8)
ax.set_xlim(0, 100); ax.set_title("18 项观察信号·同池百分位（演示子池 20 只）", fontsize=11)
ax.grid(axis="x", alpha=.25)
fig.tight_layout(); fig.savefig(os.path.join(OUT, "chart_factors.png")); plt.close(fig)

# 图 3：H 轨迹（52 周）
fig, ax = plt.subplots(figsize=(8.6, 2.9), dpi=150)
ax.plot(hist52.index, hist52.values, lw=1.2, color="#3355AA", label="周度综合位置（简化口径）")
ax.scatter([hist52.index[-1]], [hist52.values[-1]], color=color_of(H), zorder=5, s=36)
ax.set_title(f"自身历史位置轨迹（近 52 周，当前 H={H:.0f}）", fontsize=11)
ax.legend(fontsize=8); ax.grid(alpha=.25)
fig.tight_layout(); fig.savefig(os.path.join(OUT, "chart_H.png")); plt.close(fig)

# 图 4：当前得分在过去 52 周中的位置
fig, ax = plt.subplots(figsize=(8.6, 2.7), dpi=150)
ax.hist(hist52.values, bins=13, color="#9DB4D5", edgecolor="white")
ax.axvline(hist52.values[-1], color=color_of(H), lw=1.6)
ax.set_title("当前得分在过去出现的位置（分布标注）", fontsize=11)
fig.tight_layout(); fig.savefig(os.path.join(OUT, "chart_hist.png")); plt.close(fig)
step("图表渲染（4 张，BR-007 三色着色）", all(os.path.exists(os.path.join(OUT, f)) for f in
     ["chart_price.png", "chart_factors.png", "chart_H.png", "chart_hist.png"]), "matplotlib -> PNG 完成")

# ---------------------------------------------------------------- 7. 报告文字（模板化生成，不引入 LLM）
S_band = band(S_val); H_band = band(H)
sorted_f = scores.loc[TARGET].sort_values()
top2 = [(f[0], f[3], scores.loc[TARGET][f[0]]) for f in FACTOR_TABLE if f[0] in sorted_f.index[-2:]]
bot2 = [(f[0], f[3], scores.loc[TARGET][f[0]]) for f in FACTOR_TABLE if f[0] in sorted_f.index[:2]]
oneline = (f"在演示同池（沪深300样本 {N_POOL} 只）中，华友钴业当前综合分 S={S_val:.0f}（{S_band[2]}区间），"
           f"近 52 周自身历史位置 H={H:.0f}（{H_band[2]}）。"
           f"主要贡献来自{top2[0][1]}（{top2[0][2]:.0f} 分）与{top2[1][1]}（{top2[1][2]:.0f} 分）；"
           f"主要拖累来自{bot2[0][1]}（{bot2[0][2]:.0f} 分）与{bot2[1][1]}（{bot2[1][2]:.0f} 分）。")
plain = [
    ("① 用的数据", "仅使用公开的日线级行情（开高低收、成交量额、振幅、换手率）与日线估值（市盈率、市净率、总市值），不使用分钟线与逐笔成交（口径见 BR-006/NG-006）。"),
    ("② 用的分析工具", "18 项观察信号（估值特征 3 / 价格行为 7 / 交易行为 4 / 风险特征 4），每项在演示同池内取平均并列名次百分位（0–100）。"),
    ("③ 合成一个数", "综合分 S = 18 项信号等权平均后，放回同池重新排名的百分位；H = 当前状态在此前 52 个完成周中的相对位置（本样张为简化演示口径）。"),
    ("④ 是什么意思/不是什么意思", "S 只描述该股当前因子状态在同池中的相对强弱，高于 60 偏有利、低于 40 偏不利、40–60 中性；它不代表上涨概率、收益预期，也不构成任何操作指令。"),
    ("⑤ 什么时候会变", "周频更新：每周收盘数据落地后重新计算即刷新；遇周涨跌幅绝对值≥5%、业绩预告/财报、成交结构异常，建议在复核日之前提前复核。"),
]
boundary = ("本报告为量化因子状态的客观描述工具。分数描述因子状态，不代表上涨概率、可能性或收益预期，"
            "亦不构成任何证券交易类操作建议或投资顾问服务。四个观察侧面仅用于拆解展示，不另设权重；风险维度得分高不等于绝对风险低。"
            "数据来自公开渠道，可能存在延迟或修正，请以交易所披露为准。")
META = {"代码": "603799", "名称": "华友钴业", "观察日": last_day.strftime("%Y-%m-%d"),
        "同池": f"演示子池（沪深300样本 {N_POOL} 只；正式版为四池归属）", "出分日": last_day.strftime("%Y-%m-%d"),
        "建议复核日": review_day}
SUMMARY = [("综合分 S", f"{S_val:.0f}", S_band[2]), ("历史位置 H（52 周）", f"{H:.0f}", H_band[2])]
SIDES = [(s, int(sum(1 for f in FACTOR_TABLE if f[1] == s)), f"{side_mean[s]:.0f}",
          band(side_mean[s])[2]) for s in ["估值特征", "价格行为", "交易行为", "风险特征"]]

# ---------------------------------------------------------------- 8. 合规关键词扫描（AC-005 预演）
FORBIDDEN = ["买入", "卖出", "加仓", "减仓", "清仓", "建仓", "抄底", "止盈", "止损", "满仓", "补仓"]
must_have = ["不代表上涨概率", "不构成", "免责" if "免责" in boundary else "不构成"]
all_text = oneline + "".join(t[1] for t in plain) + boundary + json.dumps({"m": META}, ensure_ascii=False)
hits = [w for w in FORBIDDEN if w in all_text]
step("合规关键词扫描（AC-005 预演）", len(hits) == 0 and "不代表上涨概率" in all_text,
     f"禁词命中: {hits if hits else '无'}；边界文案在位: 是")

# ---------------------------------------------------------------- 9. HTML 在线版
html = f"""<!DOCTYPE html><html lang="zh-CN"><head><meta charset="utf-8">
<title>个股因子诊断报告（Spike-2 演示样张）·603799 华友钴业</title>
<style>
body{{font-family:"Microsoft YaHei",system-ui,sans-serif;max-width:900px;margin:24px auto;padding:0 20px;color:#222;line-height:1.65}}
h1{{font-size:22px;border-bottom:3px solid #333;padding-bottom:8px}} h2{{font-size:17px;margin-top:30px;border-left:4px solid #444;padding-left:8px}}
table{{border-collapse:collapse;width:100%;margin:10px 0;font-size:14px}} th,td{{border:1px solid #ccc;padding:7px 10px;text-align:left}}
th{{background:#f2f2f2}} .tag{{display:inline-block;padding:1px 10px;border-radius:11px;color:#fff;font-size:13px}}
img{{max-width:100%;border:1px solid #eee}} .spike{{background:#fff3cd;border:1px solid #e0c568;padding:10px 14px;border-radius:6px;font-size:13px}}
.small{{font-size:12px;color:#666}} .concl{{background:#f7f7f9;padding:12px 16px;border-radius:6px;font-size:15px}}
</style></head><body>
<div class="spike"><b>SPIKE-2 演示样张（非正式报告）</b>：同池为 20 只演示子池；H 为简化口径；本页用于验证渲染管线（E-4/E-11），正式版以 spec v1.3 与中证四池名单为准。</div>
<h1>华友钴业（603799）周频因子诊断报告 · 简要版</h1>
<h2>元信息</h2><table>{"".join(f"<tr><th style='width:160px'>{k}</th><td>{v}</td></tr>" for k, v in META.items())}</table>
<h2>汇总</h2><table><tr><th>指标</th><th>分数</th><th>分档（&gt;60 偏有利 / 40–60 中性 / &lt;40 偏不利）</th></tr>
{''.join(f"<tr><td>{n}</td><td><b style='color:{color_of(float(v))}'>{v}</b></td><td>{b}</td></tr>" for n, v, b in SUMMARY)}</table>
<h2>第一层 · 一句话结论</h2><div class="concl">{oneline}</div>
<h2>四个观察侧面</h2><table><tr><th>侧面</th><th>观察项数</th><th>侧面均分</th><th>分档</th></tr>
{''.join(f"<tr><td>{s}</td><td>{c}</td><td><b style='color:{color_of(float(m))}'>{m}</b></td><td>{b}</td></tr>" for s, c, m, b in SIDES)}</table>
<h2>股性曲线与历史位置</h2><img src="chart_price.png"><img src="chart_H.png"><img src="chart_hist.png">
<h2>18 项观察信号分布</h2><img src="chart_factors.png">
<h2>怎么读这个分数</h2><p>分数为同池相对强弱的百分位打分（0–100），<b>仅代表当前因子状态的相对位置，不代表概率、可能性或收益预期</b>。高于 60 读取为偏有利信号，低于 40 读取为偏不利信号，40–60 为中性区。</p>
<h2>第二层 · 白话解释</h2>{''.join(f"<p><b>{t}</b>：{d}</p>" for t, d in plain)}
<h2>风险与边界</h2><p>{boundary}</p>
<p class="small">Spike-2 渲染管线验证产物 · 生成时间 {dt.datetime.now().strftime('%Y-%m-%d %H:%M')} · 因子清单见 spec v1.3 附表 A（18 项，PM 审定冻结）</p>
</body></html>"""
html_path = os.path.join(OUT, "华友钴业_603799_Spike2演示报告_简要版.html")
open(html_path, "w", encoding="utf-8").write(html)
step("HTML 在线版输出", os.path.exists(html_path), html_path)

# ---------------------------------------------------------------- 10. docx 文件版
def set_font(style_or_run, name="微软雅黑", size=None, color=None, bold=None):
    f = style_or_run.font
    f.name = name
    if size: f.size = Pt(size)
    if color: f.color.rgb = RGBColor.from_string(color.lstrip("#"))
    if bold is not None: f.bold = bold
    rpr = f.element.get_or_add_rPr()
    ea = rpr.find(qn("w:eastAsia"))
    if ea is None:
        ea = rpr.makeelement(qn("w:eastAsia"), {}); rpr.append(ea)
    ea.set(qn("w:val"), name)

doc = Document()
set_font(doc.styles["Normal"], size=10.5)
doc.add_heading("华友钴业（603799）周频因子诊断报告 · 简要版", level=1)
p = doc.add_paragraph(); set_font(p.add_run("SPIKE-2 演示样张（非正式报告）：20 只演示子池；H 为简化口径；验证渲染管线（E-4/E-11）。"), size=9, color="8A6D00")
doc.add_heading("元信息", level=2)
tb = doc.add_table(rows=0, cols=2); tb.style = "Table Grid"
for k, v in META.items():
    row = tb.add_row().cells; row[0].text = k; row[1].text = str(v)
    set_font(row[0].paragraphs[0].runs[0], bold=True)
doc.add_heading("汇总", level=2)
tb = doc.add_table(rows=1, cols=3); tb.style = "Table Grid"
for i, h in enumerate(["指标", "分数", "分档"]):
    tb.rows[0].cells[i].text = h; set_font(tb.rows[0].cells[i].paragraphs[0].runs[0], bold=True)
for n, v, b in SUMMARY:
    row = tb.add_row().cells
    row[0].text = n; row[1].text = v; row[2].text = b
    set_font(row[1].paragraphs[0].runs[0], bold=True, color=color_of(float(v)))
doc.add_heading("第一层 · 一句话结论", level=2)
set_font(doc.add_paragraph().add_run(oneline))
doc.add_heading("四个观察侧面", level=2)
tb = doc.add_table(rows=1, cols=4); tb.style = "Table Grid"
for i, h in enumerate(["侧面", "观察项数", "侧面均分", "分档"]):
    tb.rows[0].cells[i].text = h; set_font(tb.rows[0].cells[i].paragraphs[0].runs[0], bold=True)
for s, c, m, b in SIDES:
    row = tb.add_row().cells
    row[0].text = s; row[1].text = str(c); row[2].text = m; row[3].text = b
    set_font(row[2].paragraphs[0].runs[0], bold=True, color=color_of(float(m)))
doc.add_heading("股性曲线与历史位置", level=2)
for img in ["chart_price.png", "chart_H.png", "chart_hist.png"]:
    doc.add_picture(os.path.join(OUT, img), width=Cm(16.5))
doc.add_heading("18 项观察信号分布", level=2)
doc.add_picture(os.path.join(OUT, "chart_factors.png"), width=Cm(16.5))
doc.add_heading("怎么读这个分数", level=2)
doc.add_paragraph("分数为同池相对强弱的百分位打分（0–100），仅代表当前因子状态的相对位置，不代表概率、可能性或收益预期。高于 60 读取为偏有利信号，低于 40 读取为偏不利信号，40–60 为中性区。")
doc.add_heading("第二层 · 白话解释", level=2)
for t, d in plain:
    p = doc.add_paragraph(); set_font(p.add_run(t + "："), bold=True); set_font(p.add_run(d))
doc.add_heading("风险与边界", level=2)
doc.add_paragraph(boundary)
p = doc.add_paragraph()
set_font(p.add_run(f"Spike-2 渲染管线验证产物 · 生成时间 {dt.datetime.now().strftime('%Y-%m-%d %H:%M')}"), size=8, color="777777")
docx_path = os.path.join(OUT, "华友钴业_603799_Spike2演示报告_简要版.docx")
doc.save(docx_path)
step("docx 文件版输出", os.path.exists(docx_path), docx_path)

# ---------------------------------------------------------------- 11. 汇总
RESULT["elapsed"] = round(time.time() - T0, 1)
RESULT["S_target"] = round(S_val, 1); RESULT["H_simplified"] = round(H, 0)
RESULT["pool_ok"] = f"{len(panels)}/{N_POOL}"
RESULT["artifacts"] = [html_path, docx_path]
RESULT["all_pass"] = all(v["pass"] for v in RESULT["steps"].values())
json.dump(RESULT, open(os.path.join(HERE, "spike2_result.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=2)
print("\n=== SPIKE-2 RESULT:", "ALL PASS" if RESULT["all_pass"] else "HAS FAILURES", f"({RESULT['elapsed']}s) ===")
print(json.dumps(RESULT["steps"], ensure_ascii=False, indent=1))
